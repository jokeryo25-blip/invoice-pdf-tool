import re
from collections.abc import Callable
import re
import json
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path
from pathlib import Path
import tkinter as tk
import tkinter as tk
from tkinter import filedialog, messagebox

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def get_application_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def get_config_path(application_dir: Path) -> Path:
    external_config = application_dir / "config.json"
    if external_config.exists() or not getattr(sys, "frozen", False):
        return external_config
    bundle_dir = Path(getattr(sys, "_MEIPASS", application_dir))
    return bundle_dir / "config.json"


def find_first_pdf(folder: Path) -> Path | None:
    pdf_files = [
        path for path in folder.iterdir() if path.is_file() and path.suffix.lower() == ".pdf"
    ]
    if not pdf_files:
        return None
    return sorted(pdf_files)[0]


def extract_text_from_pdf(pdf_path: Path) -> str:
    extracted_text = []

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text:
                extracted_text.append(page_text)

    return "\n\n".join(extracted_text)


def extract_invoice_date(pdf_text: str) -> str:
    match = re.search(r"請求日\s*:\s*(.+)", pdf_text)
    return match.group(1).strip() if match else ""


def extract_company_name(pdf_text: str) -> str:
    match = re.search(r"会社名\s*:\s*(.+)", pdf_text)
    return match.group(1).strip() if match else ""


def extract_invoice_number(pdf_text: str) -> str:
    match = re.search(r"請求番号\s*:\s*(.+)", pdf_text)
    return match.group(1).strip() if match else ""


def extract_amount(pdf_text: str) -> str:
    match = re.search(r"請求金額\s*:\s*(.+)", pdf_text)
    return match.group(1).strip() if match else ""


def append_run_log(project_dir: Path, processed_count: int, error_count: int) -> None:
    logs_dir = project_dir / "logs"
    logs_dir.mkdir(parents=True, exist_ok=True)
    log_path = logs_dir / "app.log"
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with log_path.open("a", encoding="utf-8") as log_file:
        log_file.write(
            f"実行日時: {timestamp} | 処理件数: {processed_count} | "
            f"エラー件数: {error_count}\n"
        )


def print_console_text(text: str) -> None:
    encoding = getattr(sys.stdout, "encoding", None) or "utf-8"
    safe_text = text.encode(encoding, errors="backslashreplace").decode(encoding)
    print(safe_text)


def choose_input_folder(folder_var: tk.StringVar, project_dir: Path) -> None:
    selected_folder = filedialog.askdirectory(
        initialdir=folder_var.get() or str(project_dir)
    )
    if selected_folder:
        folder_var.set(selected_folder)


def choose_output_folder(folder_var: tk.StringVar, project_dir: Path) -> None:
    selected_folder = filedialog.askdirectory(
        initialdir=folder_var.get() or str(project_dir)
    )
    if selected_folder:
        folder_var.set(selected_folder)


def show_processing_result(success_count: int, error_count: int) -> None:
    messagebox.showinfo(
        "処理結果",
        f"正常件数: {success_count}\nエラー件数: {error_count}",
    )


def main(
    input_folder: Path | None = None,
    output_folder: Path | None = None,
    on_complete: Callable[[int, int], None] | None = None,
    on_log: Callable[[str], None] | None = None,
) -> int:
    project_dir = get_application_dir()
    config_path = get_config_path(project_dir)
    with config_path.open(encoding="utf-8") as config_file:
        config = json.load(config_file)

    invoices_dir = input_folder or project_dir / config["入力フォルダ"]
    output_dir = output_folder or project_dir / config["出力フォルダ"]
    if not invoices_dir.is_absolute():
        invoices_dir = project_dir / invoices_dir
    if not output_dir.is_absolute():
        output_dir = project_dir / output_dir

    def log_status(message: str) -> None:
        if on_log:
            on_log(message)

    if not invoices_dir.exists():
        print(f"PDF folder not found: {invoices_dir}")
        log_status(f"入力フォルダが見つかりません: {invoices_dir}")
        append_run_log(project_dir, 0, 0)
        if on_complete:
            on_complete(0, 0)
        return 1

    pdf_files = sorted(
        path for path in invoices_dir.iterdir() if path.is_file() and path.suffix.lower() == ".pdf"
    )
    if not pdf_files:
        print(f"No PDF file found in: {invoices_dir}")
        log_status(f"PDFが見つかりません: {invoices_dir}")
        append_run_log(project_dir, 0, 0)
        if on_complete:
            on_complete(0, 0)
        return 1

    invoice_list = []
    invoice_file_names = []
    error_list = []
    log_status(f"PDF処理を開始します: {len(pdf_files)}件")
    for pdf_path in pdf_files:
        try:
            log_status(f"処理中: {pdf_path.name}")
            text = extract_text_from_pdf(pdf_path)
            text_length = len(text)
            print_console_text(
                f"{pdf_path.name}\n文字数: {text_length}\n先頭300文字:\n{text[:300]}"
            )
            log_status(f"{pdf_path.name} 文字数: {text_length}")
            print_console_text(f"=== {pdf_path.name} PDF全文 ===")
            print_console_text(text)
            print_console_text("==========================")
            invoice = {
                "請求日": extract_invoice_date(text),
                "会社名": extract_company_name(text),
                "請求番号": extract_invoice_number(text),
                "請求金額": extract_amount(text),
            }
            missing_fields = [field for field, value in invoice.items() if not value]
            if missing_fields:
                raise ValueError(f"抽出できない項目: {', '.join(missing_fields)}")
            invoice_list.append(invoice)
            invoice_file_names.append(pdf_path.name)
            log_status(f"処理完了: {pdf_path.name}")
        except Exception as error:
            print(f"処理をスキップしました: {pdf_path.name} ({error})")
            log_status(f"処理失敗: {pdf_path.name} ({error})")
            error_list.append(
                {
                    "ファイル名": pdf_path.name,
                    "エラー内容": str(error),
                }
            )

    append_run_log(project_dir, len(pdf_files), len(error_list))

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "請求書一覧.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Invoice"
    invoice_headers = ["請求日", "会社名", "請求番号", "請求金額"]
    headers = ["ファイル名", *invoice_headers]
    worksheet.append(headers)
    for file_name, invoice in zip(invoice_file_names, invoice_list):
        worksheet.append([file_name, *[invoice[header] for header in invoice_headers]])
        amount_cell = worksheet.cell(row=worksheet.max_row, column=headers.index("請求金額") + 1)
        amount_value = re.sub(r"[^0-9.-]", "", invoice["請求金額"])
        amount_cell.value = float(amount_value)
        amount_cell.number_format = '"¥"#,##0.##'

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for column_cells in worksheet.columns:
        max_length = max(
            sum(
                2 if unicodedata.east_asian_width(character) in {"F", "W"} else 1
                for character in str(cell.value)
            )
            for cell in column_cells
            if cell.value is not None
        )
        worksheet.column_dimensions[column_cells[0].column_letter].width = max_length + 2

    error_worksheet = workbook.create_sheet("Error")
    error_headers = ["ファイル名", "エラー内容"]
    error_worksheet.append(error_headers)
    for error in error_list:
        error_worksheet.append([error[header] for header in error_headers])

    summary_worksheet = workbook.create_sheet("Summary")
    summary_worksheet.append(["項目", "値"])
    invoice_count = worksheet.max_row - 1
    amount_column = headers.index("請求金額") + 1
    total_amount = sum(
        worksheet.cell(row=row, column=amount_column).value or 0
        for row in range(2, worksheet.max_row + 1)
    )
    average_amount = total_amount / invoice_count if invoice_count else 0
    summary_worksheet.append(["件数", invoice_count])
    summary_worksheet.append(["合計請求額", total_amount])
    summary_worksheet.append(["平均請求額", average_amount])

    workbook.save(output_path)

    print(invoice_list)
    print(f"Excel file created: {output_path}")
    log_status(f"Excel出力完了: {output_path}")
    if on_complete:
        on_complete(len(invoice_list), len(error_list))
    return 0


if __name__ == "__main__":
    root = tk.Tk()
    root.title("請求書PDF一括集計ツール")
    root.geometry("680x440")

    project_dir = get_application_dir()
    with get_config_path(project_dir).open(encoding="utf-8") as config_file:
        config = json.load(config_file)
    input_folder_var = tk.StringVar(
        value=str(project_dir / config["入力フォルダ"])
    )
    output_folder_var = tk.StringVar(
        value=str(project_dir / config["出力フォルダ"])
    )
    tk.Label(root, text="入力フォルダ").grid(row=0, column=0, padx=8, pady=12)
    tk.Entry(root, textvariable=input_folder_var).grid(
        row=0, column=1, padx=8, pady=12, sticky="ew"
    )
    tk.Button(
        root,
        text="参照",
        command=lambda: choose_input_folder(input_folder_var, project_dir),
    ).grid(row=0, column=2, padx=8, pady=12)
    tk.Label(root, text="出力フォルダ").grid(row=1, column=0, padx=8, pady=12)
    tk.Entry(root, textvariable=output_folder_var).grid(
        row=1, column=1, padx=8, pady=12, sticky="ew"
    )
    tk.Button(
        root,
        text="参照",
        command=lambda: choose_output_folder(output_folder_var, project_dir),
    ).grid(row=1, column=2, padx=8, pady=12)
    root.columnconfigure(1, weight=1)

    exit_code = 0

    tk.Label(root, text="実行ログ").grid(
        row=3, column=0, columnspan=3, padx=8, sticky="w"
    )
    log_frame = tk.Frame(root)
    log_frame.grid(row=4, column=0, columnspan=3, padx=8, pady=8, sticky="nsew")
    log_text = tk.Text(log_frame, height=10, state="disabled", wrap="word")
    log_scrollbar = tk.Scrollbar(log_frame, command=log_text.yview)
    log_text.configure(yscrollcommand=log_scrollbar.set)
    log_text.pack(side="left", fill="both", expand=True)
    log_scrollbar.pack(side="right", fill="y")
    root.rowconfigure(4, weight=1)

    def append_gui_log(message: str) -> None:
        log_text.configure(state="normal")
        log_text.insert(tk.END, f"{datetime.now():%H:%M:%S} {message}\n")
        log_text.see(tk.END)
        log_text.configure(state="disabled")
        root.update_idletasks()

    def start_processing() -> None:
        global exit_code
        input_folder = Path(input_folder_var.get()) if input_folder_var.get().strip() else None
        output_folder = Path(output_folder_var.get()) if output_folder_var.get().strip() else None
        append_gui_log("処理を開始します")
        exit_code = main(
            input_folder,
            output_folder,
            show_processing_result,
            append_gui_log,
        )

    tk.Button(root, text="請求書集計開始", command=start_processing).grid(
        row=2, column=0, columnspan=3, padx=8, pady=12
    )
    root.mainloop()
    raise SystemExit(exit_code)
